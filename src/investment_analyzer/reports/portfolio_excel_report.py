from openpyxl import Workbook
from openpyxl.styles import Font

from investment_analyzer.analysis.portfolio_analyzer import (
    PortfolioAnalyzer,
)
from investment_analyzer.analysis.scenario_analyzer import (
    ScenarioAnalyzer,
)
from investment_analyzer.analysis.historical_scenarios import (
    HISTORICAL_SCENARIOS,
)
from investment_analyzer.reports.base_report import BaseReport


class PortfolioExcelReport(BaseReport):
    """
    Creates Excel portfolio analysis reports.
    """

    def create_portfolio_report(self, portfolio):
        """
        Create an Excel workbook containing portfolio analysis.
        """

        safe_name = portfolio.name.strip().replace(" ", "_")

        if not safe_name:
            safe_name = "Portfolio"

        filename = self.report_path(f"{safe_name}_Report.xlsx")

        analyzer = PortfolioAnalyzer(portfolio)
        returns = analyzer.monthly_returns()
        growth = analyzer.growth_index(10000)

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"

        summary.append(["Investment Analyzer Portfolio Analysis Report"])
        summary["A1"].font = Font(
            bold=True,
            size=16,
        )

        summary.append([])
        summary.append(["Portfolio", portfolio.name])
        summary.append(["Generated", self.timestamp])

        summary.append([])

        summary.append(["Portfolio Holdings"])
        summary["A6"].font = Font(bold=True)

        summary.append(["Fund", "Allocation"])

        holdings_header_row = summary.max_row

        for holding in portfolio.holdings:
            summary.append(
                [
                    holding.fund.symbol,
                    holding.allocation / 100.0,
                ]
            )

        total_row = summary.max_row + 1

        summary.append(
            [
                "Total Allocation",
                portfolio.total_allocation / 100.0,
            ]
        )

        for row in range(
            holdings_header_row + 1,
            total_row + 1,
        ):
            summary.cell(
                row=row,
                column=2,
            ).number_format = "0.0%"

        summary.append([])

        summary.append(["Performance Statistics"])
        summary.cell(
            row=summary.max_row,
            column=1,
        ).font = Font(bold=True)

        first_month = returns.index.min().strftime("%b %Y")
        last_month = returns.index.max().strftime("%b %Y")

        summary.append(
            [
                "Analysis Period",
                f"{first_month} - {last_month}",
            ]
        )
        summary.append(["Months Analyzed", len(returns)])
        summary.append(["CAGR", analyzer.cagr()])
        summary.append(
            [
                "Annualized Average Return",
                analyzer.annualized_return(),
            ]
        )
        summary.append(
            [
                "Annualized Volatility",
                analyzer.annualized_volatility(),
            ]
        )
        summary.append(
            [
                "Maximum Drawdown",
                analyzer.max_drawdown(),
            ]
        )

        summary.append(["Sharpe Ratio", analyzer.sharpe_ratio()])
        summary.append(["Growth of $10,000", growth.iloc[-1]])

        percentage_labels = {
            "CAGR",
            "Annualized Average Return",
            "Annualized Volatility",
            "Maximum Drawdown",
        }

        for row in range(1, summary.max_row + 1):
            label = summary.cell(row, 1).value

            if label in percentage_labels:
                summary.cell(row, 2).number_format = "0.00%"

            elif label == "Sharpe Ratio":
                summary.cell(row, 2).number_format = "0.00"

            elif label == "Growth of $10,000":
                summary.cell(row, 2).number_format = "$#,##0.00"

        interpretation_sheet = workbook.create_sheet("Interpretation")

        interpretation_sheet.append(["Portfolio Risk & Interpretation"])
        interpretation_sheet["A1"].font = Font(
            bold=True,
            size=16,
        )

        interpretation = analyzer.interpretation()

        interpretation_sheet.append([])
        interpretation_sheet.append(["Growth", interpretation["growth"]])
        interpretation_sheet.append(["Risk", interpretation["risk"]])
        interpretation_sheet.append(["Drawdown", interpretation["drawdown"]])
        interpretation_sheet.append(
            [
                "Risk-Adjusted Performance",
                interpretation["risk_adjusted"],
            ]
        )

        interpretation_sheet.column_dimensions["A"].width = 28
        interpretation_sheet.column_dimensions["B"].width = 100
        drawdown_sheet = workbook.create_sheet("Major Drawdowns")

        drawdown_sheet.append(["Major Historical Drawdowns"])
        drawdown_sheet["A1"].font = Font(
            bold=True,
            size=16,
        )

        drawdown_sheet.append([])

        drawdown_sheet.append(
            [
                "Peak",
                "Bottom",
                "Recovery",
                "Decline",
                "Days to Bottom",
                "Days Bottom-to-Recovery",
                "Total Recovery Days",
            ]
        )

        header_row = drawdown_sheet.max_row

        for cell in drawdown_sheet[header_row]:
            cell.font = Font(bold=True)

        major_drawdowns = analyzer.major_drawdowns(threshold=0.10)

        for episode in major_drawdowns:
            drawdown_sheet.append(
                [
                    episode["peak_date"].strftime("%Y-%m"),
                    episode["bottom_date"].strftime("%Y-%m"),
                    episode["recovery_date"].strftime("%Y-%m"),
                    episode["decline"],
                    episode["days_to_bottom"],
                    episode["days_bottom_to_recovery"],
                    episode["days_to_recovery"],
                ]
            )

        for row in drawdown_sheet.iter_rows(
            min_row=header_row + 1,
            min_col=4,
            max_col=4,
        ):
            row[0].number_format = "0.00%"

        drawdown_widths = {
            "A": 14,
            "B": 14,
            "C": 14,
            "D": 14,
            "E": 18,
            "F": 25,
            "G": 20,
        }

        for column, width in drawdown_widths.items():
            drawdown_sheet.column_dimensions[column].width = width

        drawdown_sheet.freeze_panes = f"A{header_row + 1}"
        scenario_sheet = workbook.create_sheet("Historical Scenarios")

        scenario_sheet.append(["Historical Scenario Analysis"])
        scenario_sheet["A1"].font = Font(
            bold=True,
            size=16,
        )

        scenario_sheet.append([])

        scenario_sheet.append(
            [
                "Scenario",
                "Return",
                "Maximum Drawdown",
                "Peak",
                "Bottom",
                "Recovery",
                "Peak to Recovery Days",
            ]
        )

        scenario_header_row = scenario_sheet.max_row

        for cell in scenario_sheet[scenario_header_row]:
            cell.font = Font(bold=True)

        scenario_analyzer = ScenarioAnalyzer(portfolio)

        for scenario in HISTORICAL_SCENARIOS.values():
            try:
                scenario_result = scenario_analyzer.analyze_scenario(scenario)

                recovery = scenario_analyzer.recovery_analysis(
                    scenario["start_date"],
                    scenario["end_date"],
                )

            except ValueError:
                scenario_sheet.append(
                    [
                        scenario["name"],
                        None,
                        None,
                        "N/A",
                        "N/A",
                        "No data",
                        None,
                    ]
                )
                continue

            recovery_date = recovery["recovery_date"]

            scenario_sheet.append(
                [
                    scenario["name"],
                    scenario_result["total_return"],
                    scenario_result["max_drawdown"],
                    (
                        recovery["peak_date"].strftime("%Y-%m")
                        if recovery["peak_date"] is not None
                        else "N/A"
                    ),
                    (
                        recovery["bottom_date"].strftime("%Y-%m")
                        if recovery["bottom_date"] is not None
                        else "N/A"
                    ),
                    (
                        recovery_date.strftime("%Y-%m")
                        if recovery_date is not None
                        else "Not reached"
                    ),
                    recovery["days_to_recovery"],
                ]
            )

        for row in range(
            scenario_header_row + 1,
            scenario_sheet.max_row + 1,
        ):
            scenario_sheet.cell(
                row=row,
                column=2,
            ).number_format = "0.00%"

            scenario_sheet.cell(
                row=row,
                column=3,
            ).number_format = "0.00%"

        scenario_widths = {
            "A": 32,
            "B": 14,
            "C": 20,
            "D": 14,
            "E": 14,
            "F": 16,
            "G": 24,
        }

        for column, width in scenario_widths.items():
            scenario_sheet.column_dimensions[column].width = width

        scenario_sheet.freeze_panes = f"A{scenario_header_row + 1}"
        stress_sheet = workbook.create_sheet("Historical Stress")

        stress_sheet.append(["Historical Stress Interpretation"])
        stress_sheet["A1"].font = Font(
            bold=True,
            size=16,
        )

        stress_sheet.append([])

        stress_sheet.append(["Category", "Interpretation"])

        stress_header_row = stress_sheet.max_row

        for cell in stress_sheet[stress_header_row]:
            cell.font = Font(bold=True)

        stress = analyzer.stress_interpretation(threshold=0.10)

        for category, text in stress.items():
            stress_sheet.append(
                [
                    category.replace("_", " ").title(),
                    text,
                ]
            )

        stress_sheet.column_dimensions["A"].width = 28
        stress_sheet.column_dimensions["B"].width = 100

        stress_sheet.freeze_panes = f"A{stress_header_row + 1}"
        workbook.save(filename)

        return filename
